"""Write the reconciliation pack: a formatted Excel workbook and a Markdown summary."""

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NUMBER_FORMAT = "#,##0.00;(#,##0.00);-"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="999999")


def _fmt(amount):
    return f"({abs(amount):,.2f})" if amount < 0 else f"{amount:,.2f}"


def _style_table(ws, df):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        values = [str(col)] + [str(v) for v in df[col].tolist()]
        ws.column_dimensions[letter].width = min(max(len(v) for v in values) + 2, 60)
        if pd.api.types.is_numeric_dtype(df[col]) and col not in ("days_apart", "age_days"):
            for cell in ws[letter][1:]:
                cell.number_format = NUMBER_FORMAT
        if pd.api.types.is_datetime64_any_dtype(df[col]):
            for cell in ws[letter][1:]:
                cell.number_format = "DD-MMM-YYYY"
            ws.column_dimensions[letter].width = 13
    ws.freeze_panes = "A2"


def _write_summary(ws, config, statement, counts):
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws["A1"] = f"{config['company']} - Bank Reconciliation Statement"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{config['bank_account']} | As at {pd.Timestamp(config['period_end']):%d %B %Y} | {config['currency']}"
    ws["A2"].font = Font(italic=True, color="555555")

    row = 4
    for label, amount, kind in statement["lines"]:
        if kind == "balance" and row > 4:
            row += 1  # blank line before the book side
        ws.cell(row, 1, label if kind != "item" else f"    {label}")
        cell = ws.cell(row, 2, amount)
        cell.number_format = NUMBER_FORMAT
        if kind in ("balance", "total"):
            ws.cell(row, 1).font = Font(bold=True)
            cell.font = Font(bold=True)
        if kind == "total":
            for c in (ws.cell(row, 1), cell):
                c.fill = TOTAL_FILL
                c.border = Border(top=THIN, bottom=THIN)
        if kind == "check":
            row_fill = OK_FILL if statement["reconciled"] else BAD_FILL
            ws.cell(row - 0, 1).font = Font(bold=True)
            for c in (ws.cell(row, 1), cell):
                c.fill = row_fill
        row += 1

    ws.cell(row, 1, "Status").font = Font(bold=True)
    ws.cell(row, 2, "RECONCILED" if statement["reconciled"] else "NOT RECONCILED").font = Font(
        bold=True, color="006100" if statement["reconciled"] else "9C0006")

    row += 2
    ws.cell(row, 1, "Matching summary").font = Font(bold=True)
    for label, value in counts.items():
        row += 1
        ws.cell(row, 1, f"    {label}")
        ws.cell(row, 2, value)


def write_excel(path, config, results, bank_side, book_side, statement, journals, counts):
    sheets = {
        "Matched": results["matched"].drop(columns=["book_id", "bank_id"]),
        "Timing Differences": bank_side,
        "Book Adjustments": book_side,
        "Journal Entries": journals,
        "Unmatched - Bank": results["unmatched_bank"].reset_index(drop=True),
        "Unmatched - Books": results["unmatched_books"].reset_index(drop=True),
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_summary(writer.book.create_sheet("Reconciliation"), config, statement, counts)
        if "Sheet" in writer.book.sheetnames:
            del writer.book["Sheet"]
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            _style_table(writer.sheets[name], df)


def write_markdown(path, config, bank_side, book_side, statement, journals, counts):
    out = [f"# Bank Reconciliation Statement - {config['company']}", "",
           f"**{config['bank_account']}** | As at **{pd.Timestamp(config['period_end']):%d %B %Y}** "
           f"| Amounts in {config['currency']}", "",
           "| Particulars | Amount |", "|---|---:|"]
    for label, amount, kind in statement["lines"]:
        if kind == "balance" and len(out) > 6:
            out.append("| | |")
        bold = kind in ("balance", "total", "check")
        text = f"**{label}**" if bold else f"&nbsp;&nbsp;&nbsp;&nbsp;{label}"
        value = f"**{_fmt(amount)}**" if bold else _fmt(amount)
        out.append(f"| {text} | {value} |")
    status = "RECONCILED" if statement["reconciled"] else "NOT RECONCILED"
    out += ["", f"**Status: {status}**", "", "## Matching summary", ""]
    out += [f"- {label}: **{value}**" for label, value in counts.items()]

    out += ["", "## Timing differences (clear next month, no entry needed)", "",
            "| Date | Voucher | Description | Category | Amount | Age (days) |", "|---|---|---|---|---:|---:|"]
    for r in bank_side.itertuples():
        out.append(f"| {r.date:%d-%b-%Y} | {r.ref_id} | {r.description} | {r.category} | {_fmt(r.adjustment)} | {r.age_days} |")

    out += ["", "## Book adjustments (record in the cash book)", "",
            "| Date | Ref | Description | Category | Effect on books |", "|---|---|---|---|---:|"]
    for r in book_side.itertuples():
        out.append(f"| {r.date:%d-%b-%Y} | {r.ref_id} | {r.description} | {r.category} | {_fmt(r.adjustment)} |")

    out += ["", "## Suggested adjusting journal entries", "",
            "| Entry | Account | Debit | Credit |", "|---|---|---:|---:|"]
    for r in journals.itertuples():
        debit = _fmt(r.debit) if pd.notna(r.debit) else ""
        credit = _fmt(r.credit) if pd.notna(r.credit) else ""
        out.append(f"| {r.entry} | {r.account.strip() if credit else r.account} | {debit} | {credit} |")

    with open(path, "w") as f:
        f.write("\n".join(out) + "\n")


def print_summary(config, statement, counts):
    width = 62
    print("=" * width)
    print(f"{config['company']} - Bank Reconciliation")
    print(f"As at {pd.Timestamp(config['period_end']):%d %B %Y} ({config['currency']})")
    print("=" * width)
    for label, amount, kind in statement["lines"]:
        if kind == "balance":
            print()
        indent = "  " if kind == "item" else ""
        print(f"{indent}{label:<{50 - len(indent)}}{_fmt(amount):>14}")
        if kind == "total":
            print("-" * width)
    print()
    for label, value in counts.items():
        print(f"  {label:<48}{value:>14}")
    print("=" * width)
    print("STATUS:", "RECONCILED" if statement["reconciled"] else "NOT RECONCILED - review exceptions")
